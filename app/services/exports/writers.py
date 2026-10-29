import csv
import io
import re
from collections.abc import Iterator, Sequence
from typing import Final
from xml.sax.saxutils import escape

from openpyxl import Workbook
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Flowable, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from app.schemas.exports import ExportFormat
from app.services.exports.tables import Cell, Document, Row
from app.services.exports.tables import Table as DataTable

FORMULA_PREFIXES: Final = ("=", "+", "-", "@", "\t", "\r")
MEDIA_TYPES: Final = {
    ExportFormat.CSV: "text/csv; charset=utf-8",
    ExportFormat.XLSX: "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    ExportFormat.PDF: "application/pdf",
}
SHEET_TITLE_MAX: Final = 31
SHEET_TITLE_FORBIDDEN: Final = re.compile(r"[\[\]:*?/\\]")
PDF_MARGIN: Final = 12 * mm
PDF_FONT_SIZE: Final = 6.5
PDF_HEADER_FILL: Final = colors.HexColor("#E8EAF0")
PDF_GRID: Final = colors.HexColor("#B8BCC8")


def safe_cell(value: Cell) -> Cell:
    if isinstance(value, str) and value.startswith(FORMULA_PREFIXES):
        return f"'{value}"
    return value


def _safe_row(row: Sequence[Cell]) -> list[Cell]:
    return [safe_cell(value) for value in row]


def _header_rows(document: Document) -> list[Row]:
    return [(label, value) for label, value in document.header]


def csv_chunks(document: Document) -> Iterator[bytes]:
    buffer = io.StringIO()
    writer = csv.writer(buffer, lineterminator="\r\n")

    def flush() -> bytes:
        chunk = buffer.getvalue().encode("utf-8")
        buffer.seek(0)
        buffer.truncate(0)
        return chunk

    for row in _header_rows(document):
        writer.writerow(_safe_row(row))
    for table in document.tables:
        writer.writerow([])
        if len(document.tables) > 1:
            writer.writerow(_safe_row((table.title,)))
        writer.writerow(_safe_row(table.columns))
        yield flush()
        for row in table.rows:
            writer.writerow(_safe_row(row))
            yield flush()
    yield flush()


def sheet_title(title: str, used: set[str]) -> str:
    base = SHEET_TITLE_FORBIDDEN.sub(" ", title)[:SHEET_TITLE_MAX].strip() or "Sheet"
    candidate = base
    number = 2
    while candidate in used:
        suffix = f" {number}"
        candidate = f"{base[: SHEET_TITLE_MAX - len(suffix)]}{suffix}"
        number += 1
    used.add(candidate)
    return candidate


def xlsx_bytes(document: Document) -> bytes:
    workbook = Workbook(write_only=True)
    used: set[str] = set()
    for index, table in enumerate(document.tables):
        sheet = workbook.create_sheet(sheet_title(table.title, used))
        if index == 0:
            for row in _header_rows(document):
                sheet.append(_safe_row(row))
            sheet.append([])
        sheet.append(_safe_row(table.columns))
        for row in table.rows:
            sheet.append(_safe_row(row))
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def _pdf_cell(value: Cell, style: ParagraphStyle) -> Paragraph:
    text = "" if value is None else str(value)
    return Paragraph(escape(text), style)


def _pdf_table(table: DataTable, width: float, style: ParagraphStyle) -> Table:
    column_width = width / max(len(table.columns), 1)
    data = [[_pdf_cell(value, style) for value in table.columns]]
    data.extend([_pdf_cell(value, style) for value in row] for row in table.rows)
    result = Table(data, colWidths=[column_width] * len(table.columns), repeatRows=1)
    result.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), PDF_HEADER_FILL),
                ("GRID", (0, 0), (-1, -1), 0.25, PDF_GRID),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ]
        )
    )
    return result


def pdf_bytes(document: Document) -> bytes:
    output = io.BytesIO()
    page = landscape(A4)
    template = SimpleDocTemplate(
        output,
        pagesize=page,
        leftMargin=PDF_MARGIN,
        rightMargin=PDF_MARGIN,
        topMargin=PDF_MARGIN,
        bottomMargin=PDF_MARGIN,
        title=f"Ibotix {document.report} report",
    )
    styles = getSampleStyleSheet()
    cell_style = ParagraphStyle(
        "cell", parent=styles["BodyText"], fontSize=PDF_FONT_SIZE, leading=PDF_FONT_SIZE + 1.5
    )
    story: list[Flowable] = [Paragraph(escape(f"Ibotix {document.report} report"), styles["Title"])]
    story.extend(
        Paragraph(f"<b>{escape(label)}:</b> {escape(value)}", styles["BodyText"])
        for label, value in document.header
    )
    width = page[0] - 2 * PDF_MARGIN
    for table in document.tables:
        story.append(Spacer(1, 4 * mm))
        if len(document.tables) > 1:
            story.append(Paragraph(escape(table.title), styles["Heading3"]))
        story.append(_pdf_table(table, width, cell_style))
    template.build(story)
    return output.getvalue()


def render(document: Document, export_format: ExportFormat) -> Iterator[bytes]:
    if export_format is ExportFormat.CSV:
        return csv_chunks(document)
    if export_format is ExportFormat.XLSX:
        return iter((xlsx_bytes(document),))
    return iter((pdf_bytes(document),))
